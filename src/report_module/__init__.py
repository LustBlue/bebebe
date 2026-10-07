"""Модуль формирования отчётов (ReportModule).

Содержит функции для формирования отчётов по задачам
(по исполнителю, по проекту, за период) и экспорта в PDF/Excel.
"""

from __future__ import annotations

import json
from datetime import datetime
from typing import Any

from src.core.config import settings
from src.core.database import get_connection, transaction
from src.core.errors import ValidationError
from src.report_module.pdf_writer import PdfBuilder, write_pdf


def generate_report_by_user(
    user_id: int,
    start_date: str | None = None,
    end_date: str | None = None,
    generated_by: int | None = None
) -> dict[str, Any]:
    """Сформировать отчёт по исполнителю.

    Args:
        user_id: ID исполнителя
        start_date: Дата начала периода (ISO формат)
        end_date: Дата окончания периода (ISO формат)
        generated_by: ID пользователя, создавшего отчёт

    Returns:
        Словарь с данными отчёта
    """
    conn = get_connection()

    # Проверка существования пользователя
    user = conn.execute(
        "SELECT id, username, email, role FROM users WHERE id = ?",
        (user_id,)
    ).fetchone()

    if not user:
        raise ValidationError("Пользователь не найден")

    # Построение запроса с фильтрами по датам
    query = """
        SELECT
            t.id, t.title, t.description, t.status, t.priority,
            t.due_date, t.created_at, t.updated_at,
            p.name as project_name
        FROM tasks t
        JOIN projects p ON t.project_id = p.id
        WHERE t.assigned_to = ?
    """

    params = [user_id]

    if start_date:
        query += " AND t.created_at >= ?"
        params.append(start_date)

    if end_date:
        query += " AND t.created_at <= ?"
        params.append(end_date)

    query += " ORDER BY t.created_at DESC"

    tasks = conn.execute(query, params).fetchall()
    tasks_list = [dict(task) for task in tasks]

    # Статистика
    total_tasks = len(tasks_list)
    status_counts = {}
    priority_counts = {}

    for task in tasks_list:
        status = task['status']
        priority = task['priority']

        status_counts[status] = status_counts.get(status, 0) + 1
        priority_counts[priority] = priority_counts.get(priority, 0) + 1

    # Сохранение отчёта
    report_data = {
        'report_type': 'by_user',
        'user': dict(user),
        'period': {
            'start_date': start_date,
            'end_date': end_date
        },
        'statistics': {
            'total_tasks': total_tasks,
            'by_status': status_counts,
            'by_priority': priority_counts
        },
        'tasks': tasks_list
    }

    if generated_by:
        report_id = _save_report('by_user', json.dumps(report_data), generated_by)
        report_data['report_id'] = report_id

    return report_data


def generate_report_by_project(
    project_id: int,
    start_date: str | None = None,
    end_date: str | None = None,
    generated_by: int | None = None
) -> dict[str, Any]:
    """Сформировать отчёт по проекту.

    Args:
        project_id: ID проекта
        start_date: Дата начала периода (ISO формат)
        end_date: Дата окончания периода (ISO формат)
        generated_by: ID пользователя, создавшего отчёт

    Returns:
        Словарь с данными отчёта
    """
    conn = get_connection()

    # Проверка существования проекта
    project = conn.execute(
        "SELECT id, name, description, created_at FROM projects WHERE id = ?",
        (project_id,)
    ).fetchone()

    if not project:
        raise ValidationError("Проект не найден")

    # Построение запроса
    query = """
        SELECT
            t.id, t.title, t.description, t.status, t.priority,
            t.due_date, t.created_at, t.updated_at,
            u1.username as assignee_name,
            u2.username as creator_name
        FROM tasks t
        LEFT JOIN users u1 ON t.assigned_to = u1.id
        JOIN users u2 ON t.created_by = u2.id
        WHERE t.project_id = ?
    """

    params = [project_id]

    if start_date:
        query += " AND t.created_at >= ?"
        params.append(start_date)

    if end_date:
        query += " AND t.created_at <= ?"
        params.append(end_date)

    query += " ORDER BY t.created_at DESC"

    tasks = conn.execute(query, params).fetchall()
    tasks_list = [dict(task) for task in tasks]

    # Статистика
    total_tasks = len(tasks_list)
    status_counts = {}
    assignee_counts = {}

    for task in tasks_list:
        status = task['status']
        assignee = task['assignee_name'] or 'Не назначено'

        status_counts[status] = status_counts.get(status, 0) + 1
        assignee_counts[assignee] = assignee_counts.get(assignee, 0) + 1

    # Сохранение отчёта
    report_data = {
        'report_type': 'by_project',
        'project': dict(project),
        'period': {
            'start_date': start_date,
            'end_date': end_date
        },
        'statistics': {
            'total_tasks': total_tasks,
            'by_status': status_counts,
            'by_assignee': assignee_counts
        },
        'tasks': tasks_list
    }

    if generated_by:
        report_id = _save_report('by_project', json.dumps(report_data), generated_by)
        report_data['report_id'] = report_id

    return report_data


def generate_report_by_period(
    start_date: str,
    end_date: str,
    generated_by: int | None = None
) -> dict[str, Any]:
    """Сформировать отчёт за период.

    Args:
        start_date: Дата начала периода (ISO формат)
        end_date: Дата окончания периода (ISO формат)
        generated_by: ID пользователя, создавшего отчёт

    Returns:
        Словарь с данными отчёта
    """
    conn = get_connection()

    # Запрос всех задач за период
    query = """
        SELECT
            t.id, t.title, t.description, t.status, t.priority,
            t.due_date, t.created_at, t.updated_at,
            p.name as project_name,
            u1.username as assignee_name,
            u2.username as creator_name
        FROM tasks t
        JOIN projects p ON t.project_id = p.id
        LEFT JOIN users u1 ON t.assigned_to = u1.id
        JOIN users u2 ON t.created_by = u2.id
        WHERE t.created_at >= ? AND t.created_at <= ?
        ORDER BY t.created_at DESC
    """

    tasks = conn.execute(query, (start_date, end_date)).fetchall()
    tasks_list = [dict(task) for task in tasks]

    # Статистика
    total_tasks = len(tasks_list)
    status_counts = {}
    project_counts = {}
    priority_counts = {}

    for task in tasks_list:
        status = task['status']
        project = task['project_name']
        priority = task['priority']

        status_counts[status] = status_counts.get(status, 0) + 1
        project_counts[project] = project_counts.get(project, 0) + 1
        priority_counts[priority] = priority_counts.get(priority, 0) + 1

    # Сохранение отчёта
    report_data = {
        'report_type': 'by_period',
        'period': {
            'start_date': start_date,
            'end_date': end_date
        },
        'statistics': {
            'total_tasks': total_tasks,
            'by_status': status_counts,
            'by_project': project_counts,
            'by_priority': priority_counts
        },
        'tasks': tasks_list
    }

    if generated_by:
        report_id = _save_report('by_period', json.dumps(report_data), generated_by)
        report_data['report_id'] = report_id

    return report_data


def export_report_to_excel(report_data: dict[str, Any]) -> str:
    """Экспортировать отчёт в Excel.

    Args:
        report_data: Данные отчёта

    Returns:
        Путь к созданному файлу
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as error:
        raise ValidationError(
            "Библиотека openpyxl не установлена: экспорт в Excel недоступен"
        ) from error

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    # Заголовок
    ws['A1'] = f"Отчёт: {report_data['report_type']}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:E1')

    # Период
    row = 3
    if 'period' in report_data:
        period = report_data['period']
        ws[f'A{row}'] = "Период:"
        ws[f'B{row}'] = f"{period['start_date']} - {period['end_date']}"
        row += 2

    # Статистика
    ws[f'A{row}'] = "Статистика"
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    stats = report_data.get('statistics', {})
    ws[f'A{row}'] = "Всего задач:"
    ws[f'B{row}'] = stats.get('total_tasks', 0)
    row += 2

    # Заголовки таблицы задач
    headers = ['ID', 'Название', 'Статус', 'Приоритет', 'Создано']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')

    row += 1

    # Данные задач
    for task in report_data.get('tasks', []):
        ws.cell(row=row, column=1, value=task['id'])
        ws.cell(row=row, column=2, value=task['title'])
        ws.cell(row=row, column=3, value=task['status'])
        ws.cell(row=row, column=4, value=task['priority'])
        ws.cell(row=row, column=5, value=task['created_at'])
        row += 1

    # Сохранение файла
    filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = settings.reports_dir / filename
    wb.save(filepath)

    return str(filepath)


def export_report_to_pdf(report_data: dict[str, Any]) -> str:
    """Экспортировать отчёт в PDF.

    Используется собственный генератор :mod:`src.report_module.pdf_writer`,
    не требующий внешних библиотек (``reportlab``), что позволяет
    формировать отчёты на изолированных рабочих местах.

    Args:
        report_data: Данные отчёта

    Returns:
        Путь к созданному файлу

    Raises:
        ValidationError: Если структура отчёта некорректна
    """
    if not isinstance(report_data, dict) or 'report_type' not in report_data:
        raise ValidationError("Некорректные данные отчёта для экспорта в PDF")

    filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.pdf"
    filepath = settings.reports_dir / filename

    builder = PdfBuilder(title=f"Отчёт ProjectFlow: {report_data['report_type']}")
    titles = {
        'by_user': 'Отчёт по исполнителю',
        'by_project': 'Отчёт по проекту',
        'by_period': 'Отчёт за период',
    }
    builder.add_heading(titles.get(report_data['report_type'],
                                   f"Отчёт: {report_data['report_type']}"))

    period = report_data.get('period') or {}
    if period:
        builder.add_key_value(
            "Период",
            f"{period.get('start_date') or 'не ограничен'} — "
            f"{period.get('end_date') or 'не ограничен'}",
        )

    for key in ('user', 'project'):
        entity = report_data.get(key)
        if isinstance(entity, dict):
            builder.add_key_value("Объект отчёта", str(entity.get('name')
                                                        or entity.get('username') or ''))
            if entity.get('email'):
                builder.add_key_value("Email", str(entity['email']))

    builder.add_heading("Сводная статистика", size=12.0, space_before=6.0)
    statistics = report_data.get('statistics') or {}
    builder.add_key_value("Всего задач", str(statistics.get('total_tasks', 0)))
    for key, label in (('by_status', 'По статусам'),
                       ('by_priority', 'По приоритетам'),
                       ('by_assignee', 'По исполнителям'),
                       ('by_project', 'По проектам')):
        values = statistics.get(key)
        if isinstance(values, dict) and values:
            joined = ", ".join(f"{name}: {count}" for name, count in sorted(values.items()))
            builder.add_key_value(label, joined)

    tasks = report_data.get('tasks') or []
    builder.add_heading("Перечень задач", size=12.0, space_before=6.0)
    rows = [
        (
            task.get('id', ''),
            task.get('title', ''),
            task.get('status', ''),
            task.get('priority', ''),
            task.get('assignee_name') or task.get('assigned_to') or '—',
            task.get('due_date') or '—',
        )
        for task in tasks
    ]
    builder.add_table(
        ["ID", "Название", "Статус", "Приоритет", "Исполнитель", "Срок"],
        rows,
        widths=[0.6, 3.2, 1.6, 1.4, 1.6, 1.4],
    )

    if not builder.lines:
        builder.add_paragraph("Отчёт не содержит данных.")

    written_path, warnings = write_pdf(filepath, builder=builder)
    if warnings:
        # Предупреждения о замене символов не прерывают экспорт.
        print("[PDF] " + "; ".join(sorted(set(warnings))))

    return str(written_path)


def _save_report(report_type: str, params: str, generated_by: int) -> int:
    """Сохранить информацию об отчёте в базу данных.

    Args:
        report_type: Тип отчёта
        params: Параметры отчёта (JSON)
        generated_by: ID создателя

    Returns:
        ID созданной записи
    """
    with transaction() as conn:
        cursor = conn.execute(
            """
            INSERT INTO reports (name, type, params, created_by)
            VALUES (?, ?, ?, ?)
            """,
            (f"Отчёт {report_type}", report_type, params, generated_by)
        )
        return cursor.lastrowid


def get_reports_list(
    user_id: int | None = None,
    limit: int = 50
) -> list[dict[str, Any]]:
    """Получить список отчётов.

    Args:
        user_id: Фильтр по создателю (опционально)
        limit: Максимальное количество отчётов

    Returns:
        Список отчётов
    """
    conn = get_connection()

    query = """
        SELECT
            r.id, r.name, r.type, r.created_at,
            u.username as created_by_name
        FROM reports r
        LEFT JOIN users u ON r.created_by = u.id
    """

    params = []

    if user_id:
        query += " WHERE r.created_by = ?"
        params.append(user_id)

    query += " ORDER BY r.created_at DESC LIMIT ?"
    params.append(limit)

    reports = conn.execute(query, params).fetchall()

    return [dict(report) for report in reports]
